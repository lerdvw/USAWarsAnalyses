# -*- coding: utf-8 -*-
"""
Sanity checks for the dataset and everything built from it.

    python3 build.py && python3 check.py

Prints a summary, then every failed check; exits with status 1 if any failed.
Outputs not built yet are skipped.
"""

import json
import os
import re

import data
from data import ROWS

HERE = os.path.dirname(os.path.abspath(__file__))
FAILED = []


def check(ok, message):
    """Record a failure; the run carries on, so every failure gets listed."""
    if not ok:
        FAILED.append(message)


def total(key):
    """A field summed over every conflict, skipping blanks."""
    return sum(row[key] for row in ROWS if row[key] is not None)


# ---------------------------------------------------------------------------
# The dataset
# ---------------------------------------------------------------------------

def rows():
    """Numbering, dates, day counts, and casualties that add up."""
    check([row["idx"] for row in ROWS] == list(range(1, len(ROWS) + 1)), "# does not run 1..n")
    for before, after in zip(ROWS, ROWS[1:]):
        check(before["start"] <= after["start"], f"{after['name']} starts before {before['name']}, above it")
    for row in ROWS:
        name = row["name"]
        check(row["end"] is None or row["end"] >= row["start"], f"{name}: ends before it starts")
        check(0 <= row["combat_days"] <= row["all_days"], f"{name}: more combat days than conflict days")
        if row["casualties"] is not None:
            check(row["casualties"] == (row["deaths"] or 0) + (row["wounded"] or 0),
                  f"{name}: casualties are not deaths plus wounded")
    ongoing = sum(row["ongoing"] for row in ROWS)
    print(f"rows            {len(ROWS)}, {ongoing} ongoing")


def scales():
    """Every authorization level has a label; print how many conflicts have each."""
    for row in ROWS:
        check(row["auth_level"] in data.LV_LABEL, f"{row['name']}: unknown authorization level {row['auth_level']}")
    counts = [f"{data.LV_LABEL[level]} {sum(row['auth_level'] == level for row in ROWS)}"
              for level in sorted(data.LV_LABEL, reverse=True)]
    print(f"authorization   {', '.join(counts)}")


def sources():
    """Sources are numbered 1..n, linked securely, and every citation exists."""
    numbers = [number for number, label, url in data.REFS]
    check(numbers == list(range(1, len(numbers) + 1)), "REFS is not numbered 1..n")
    for number, label, url in data.REFS:
        check(url is None or url.startswith("https://"), f"source {number} is not an https link")
    for row in ROWS:
        check(row["sources"] and all(1 <= n <= len(numbers) for n in row["sources"]),
              f"{row['name']}: cites a source that is not in REFS")
    cited = {n for row in ROWS for n in row["sources"]}
    print(f"sources         {len(numbers)} listed, {len(cited)} cited")


def totals():
    """Print the headline totals and how many rows lack a figure."""
    blank = sum(row["kia"] is None or row["cost_m"] is None for row in ROWS)
    print(f"KIA + MIA       {total('kia'):,}")
    print(f"casualties      {total('casualties'):,}")
    print(f"cost, then-year ${total('cost_m'):,}m")
    print(f"blank figures   {blank} rows lack a death count or a cost")


def inflation():
    """No cost is dated to a year dearer than the base month."""
    for row in ROWS:
        check(row["factor"] >= 1, f"{row['name']}: {row['cost_year']} prices sit above the base month")
    print(f"cost, {data.CPI_BASE_LABEL} ${total('cost_adj'):,.0f}m")


def tallies():
    """The Apache row carries the sums of its engagement table, not typed-in numbers."""
    table = data.APACHE_ENGAGEMENTS
    check(all(low <= high for _, _, low, high, _, _ in table), "an Apache engagement's low exceeds its high")
    check([e[0] for e in table] == sorted(e[0] for e in table), "Apache engagements are out of date order")
    row = next(row for row in ROWS if row["name"] == "Apache Wars, final phase")
    check(data.APACHE_KILLED_LOW <= row["kia"] <= data.APACHE_KILLED_HIGH, "Apache KIA falls outside its range")
    check(row["wounded"] == data.APACHE_WOUNDED, "Apache wounded differ from the engagement table")
    print(f"apache tally    {len(table)} engagements, killed {data.APACHE_KILLED_LOW}-{data.APACHE_KILLED_HIGH} "
          f"(row {row['kia']}), wounded {data.APACHE_WOUNDED}")


def crs():
    """Compare CRS's constant-dollar figures with the CPI method, war by war."""
    check(set(data.CRS_FY2011) <= {row["name"] for row in ROWS}, "CRS_FY2011 names a conflict not in ROWS")
    for row in ROWS:
        if row["crs_2011"] and row["cost_adj"]:
            ratio = row["crs_2011"] * data.CPI_BASE / 224.939 / row["cost_adj"]
            print(f"  CRS / CPI {ratio:5.2f}  {row['name']}")


# ---------------------------------------------------------------------------
# What build.py makes
# ---------------------------------------------------------------------------

def figures():
    """The Figures workbook holds every conflict in its own row, in order."""
    path = os.path.join(HERE, "US-Conflicts-1775-2026-Figures.xlsx")
    if not os.path.exists(path):
        print("figures.xlsx    not built; run build.py first")
        return
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(path)
    ws = wb["Figures"]
    first, last = 6, 5 + len(ROWS)
    headers = [cell.value for cell in ws[5] if cell.value]
    column = {header: get_column_letter(n) for n, header in enumerate(headers, 1)}
    check(ws.tables["Figures"].ref == f"A5:{get_column_letter(len(headers))}{last}",
          "the Figures table does not cover every row and column")
    # Rows run from the largest August 2026 cost down; those without one follow in date order.
    order = sorted(ROWS, key=lambda r: (r["cost_adj"] is None, -(r["cost_adj"] or 0), r["idx"]))
    at = {row["name"]: 6 + n for n, row in enumerate(order)}
    for row in ROWS:
        name = ws[f"{column['Conflict']}{at[row['name']]}"].value.replace(" ‡", "")
        check(name == row["name"], f"Figures row {at[row['name']]} holds {name}, not {row['name']}")
    costs = [row["cost_adj"] for row in order if row["cost_adj"] is not None]
    check(costs == sorted(costs, reverse=True), "Figures rows are not in falling cost order")
    # Each total sums every conflict's row.
    for header in headers:
        if header in ("Combat Days", "All Conflict Days", "KIA + MIA", "All Casualties") or header.startswith("Net Cost"):
            letter = column[header]
            check(ws[f"{letter}{last + 1}"].value == f"=SUM({letter}{first}:{letter}{last})",
                  f"the {header} total misses rows")
    # Each conflict is priced by its own row of the deflator sheet.
    deflator = wb["CPI_Deflator"]
    cost = column[next(h for h in headers if h.startswith("Net Cost to"))]
    adjusted = column[next(h for h in headers if h.startswith("Net Cost in"))]
    for row in ROWS:
        r = at[row["name"]]
        check(deflator[f"B{r}"].value == row["name"], f"CPI_Deflator row {r} is not {row['name']}")
        check(ws[f"{adjusted}{r}"].value == f'=IF({cost}{r}="","",{cost}{r}*CPI_Deflator!F{r})',
              f"Figures row {r} is priced from another conflict's deflator row")
    # The text columns carry the dataset's words, and each row its sources.
    for row in ROWS:
        r = at[row["name"]]
        for header, key in (("What Happened", "summary"), ("Casualty Detail", "losses_text"),
                            ("Cost Accounting", "cost_text")):
            check(ws[f"{column[header]}{r}"].value == row[key], f"Figures row {r}: {header} differs from data.py")
        check(ws[f"{column['Sources']}{r}"].value == ", ".join(str(n) for n in row["sources"]),
              f"Figures row {r}: sources differ from data.py")
    # Authorization reads exactly as it does in the Record workbook.
    for row in ROWS:
        r = at[row["name"]]
        check(ws[f"{column['Authorization']}{r}"].value == f"{row['auth_label']} - {row['auth_note']}",
              f"Figures row {r}: authorization differs from the Record")
    print(f"figures.xlsx    {len(wb.sheetnames)} sheets, {last - first + 1} rows")


def record():
    """The Record workbook holds every conflict in its own row, in order."""
    path = os.path.join(HERE, "US-Conflicts-1775-2026-Record.xlsx")
    if not os.path.exists(path):
        print("record.xlsx     not built; run build.py first")
        return
    from openpyxl import load_workbook
    ws = load_workbook(path)["Record"]
    check(ws.tables["Record"].ref == f"A5:L{5 + len(ROWS)}", "the Record table does not cover every row")
    for row in ROWS:
        name = ws.cell(row=5 + row["idx"], column=2).value.split("\n")[0]
        check(name == row["name"], f"Record row {5 + row['idx']} holds {name}, not {row['name']}")
    print(f"record.xlsx     {len(ROWS)} rows")


def json_block(text, block_id):
    """Parse one <script type="application/json" id="..."> block of a page."""
    match = re.search(r'<script type="application/json" id="' + block_id + r'">\s*(.*?)\s*</script>', text, re.S)
    return json.loads(match.group(1)) if match else []


def pages():
    """Both pages carry every conflict, in order."""
    columns = {}
    for name in ("figures.html", "record.html"):
        path = os.path.join(HERE, name)
        if not os.path.exists(path):
            print(f"{name:<15} not built; run build.py first")
            continue
        with open(path, encoding="utf-8") as f:
            text = f.read()
        rows = json_block(text, "table-rows")
        columns[name] = {c["key"]: c for c in json_block(text, "table-columns")}
        check([r["name"] for r in rows] == [row["name"] for row in ROWS], f"{name} does not carry every conflict in order")
        print(f"{name:<15} {len(rows)} rows, {len(text):,} bytes")
    if "figures.html" in columns:
        check({"summary", "losses_text", "cost_text"} <= set(columns["figures.html"]),
              "figures.html lacks What happened, Casualty detail or Cost accounting")
    if len(columns) == 2:
        check(columns["figures.html"]["auth_label"] == columns["record.html"]["auth_label"],
              "the two pages show authorization differently")


if __name__ == "__main__":
    rows()
    scales()
    sources()
    totals()
    inflation()
    tallies()
    crs()
    figures()
    record()
    pages()
    for message in FAILED:
        print("FAIL", message)
    print("\nall checks pass" if not FAILED else f"\n{len(FAILED)} check(s) failed")
    raise SystemExit(1 if FAILED else 0)
